#!/usr/bin/env python3
"""Builds the per-sample metagenomics_report.pdf (via generate_report_v3's
generate_pdf(), CNV plot + host-depletion stats embedded) and a companion
summary.json, in one pass so both stay consistent.
"""
import argparse
import json
import os
import re
import sys
from datetime import date


def parse_flagstat(path):
    text = open(path).read()
    primary_total = primary_mapped = None
    for line in text.splitlines():
        m = re.match(r"^(\d+)\s*\+\s*\d+\s+primary$", line.strip())
        if m:
            primary_total = int(m.group(1))
        m = re.match(r"^(\d+)\s*\+\s*\d+\s+primary mapped", line.strip())
        if m:
            primary_mapped = int(m.group(1))
    if primary_total is None:
        m = re.search(r"^(\d+)\s*\+\s*\d+\s+in total", text, re.MULTILINE)
        primary_total = int(m.group(1)) if m else 0
        m = re.search(r"^(\d+)\s*\+\s*\d+\s+mapped", text, re.MULTILINE)
        primary_mapped = int(m.group(1)) if m else 0
    non_human = primary_total - primary_mapped
    return {
        "total_reads": primary_total,
        "human_reads": primary_mapped,
        "non_human_reads": non_human,
        "pct_human": round(100 * primary_mapped / primary_total, 2) if primary_total else 0.0,
        "pct_non_human": round(100 * non_human / primary_total, 2) if primary_total else 0.0,
    }


def parse_ichorcna_params(path):
    result = {}
    with open(path) as fh:
        for line in fh:
            if "\t" not in line:
                continue
            key, val = line.rstrip("\n").split("\t", 1)
            result[key.rstrip(":").strip()] = val.strip()
    return {
        "tumor_fraction": result.get("Tumor Fraction"),
        "ploidy": result.get("Ploidy"),
        "gender": result.get("Gender"),
        "subclone_fraction": result.get("Subclone Fraction"),
        "fraction_genome_subclonal": result.get("Fraction Genome Subclonal"),
        "fraction_cna_subclonal": result.get("Fraction CNA Subclonal"),
        "coverage": result.get("Coverage"),
    }


def format_mio_de(n):
    """z.B. 2734281 -> '2,7' (deutsche Dezimalschreibweise, eine Nachkommastelle)."""
    return f"{n / 1_000_000:.1f}".replace(".", ",")


NEXUS_BEFUND_TEXT = (
    "Aus FFPE-Material wurde DNA isoliert und mittels Nanopore-Sequenzierung "
    "untersucht. Die Rohdaten wurden mit Kraken-Uniq analysiert (Breitwieser "
    "et al. 2018). Aus den humanen Reads wurde mit ichorCNA ein CNV-Profil "
    "erstellt (Adalsteinsson et al. 2017). Bei insgesamt {mio} Mio. Reads "
    "fanden sich keine relevanten Erreger."
)
NEXUS_BEFUND_TEXT_2 = "Im CNV-Profil fanden sich keine relevanten Alterationen (siehe Abbildung)."


def build_nexus_befund_docx(output_path, total_reads, cnv_plot_path):
    """Word-Vorlagenbefund (Calibri 11pt) mit eingefügtem CNV-Profil.
    Fester Textbaustein für den unauffälligen Befund -- bei tatsächlichen
    Erreger-/CNV-Befunden ist der Text vor Ausgabe manuell anzupassen."""
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def add_paragraph(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        return p

    add_paragraph(NEXUS_BEFUND_TEXT.format(mio=format_mio_de(total_reads)))
    add_paragraph(NEXUS_BEFUND_TEXT_2)

    if cnv_plot_path and os.path.exists(cnv_plot_path):
        doc.add_picture(cnv_plot_path, width=Inches(6.3))

    doc.save(output_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sample", required=True)
    ap.add_argument("--platform", required=True, choices=["nanopore", "illumina"])
    ap.add_argument("--flagstat", required=True)
    ap.add_argument("--ichorcna-params", required=True)
    ap.add_argument("--ichorcna-plot", required=True)
    ap.add_argument("--krakenuniq-report", default="")
    ap.add_argument("--output-pdf", required=True)
    ap.add_argument("--output-json", required=True)
    ap.add_argument("--output-docx", required=True)
    args = ap.parse_args()

    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from generate_report_v3 import generate_pdf, parse_report, select_top_hits  # noqa: E402

    os.makedirs(os.path.dirname(args.output_pdf) or ".", exist_ok=True)

    host = parse_flagstat(args.flagstat)
    cnv = parse_ichorcna_params(args.ichorcna_params)
    cnv["plot_path"] = args.ichorcna_plot
    extra = {"host": host, "cnv": cnv}

    krakenuniq_ran = bool(args.krakenuniq_report) and os.path.exists(args.krakenuniq_report)
    report_text = open(args.krakenuniq_report).read() if krakenuniq_ran else None
    today = date.today().strftime("%d. %B %Y").lstrip("0")

    generate_pdf(
        output_path=args.output_pdf,
        report_text=report_text,
        sample_name=args.sample,
        date=today,
        krakenuniq_ran=krakenuniq_ran,
        extra=extra,
    )

    os.makedirs(os.path.dirname(args.output_docx) or ".", exist_ok=True)
    build_nexus_befund_docx(args.output_docx, host["total_reads"], args.ichorcna_plot)

    krakenuniq_json = None
    if krakenuniq_ran:
        parsed = parse_report(report_text)
        top_hits, _rare, _common = select_top_hits(parsed["hits"])
        krakenuniq_json = {
            "total_all": parsed["total_all"],
            "total_classified": parsed["total_classified"],
            "total_unclassified": parsed["total_unclassified"],
            "pct_classified": parsed["pct_classified"],
            "pct_unclassified": parsed["pct_unclassified"],
            "top_hits": [
                {"name": h.get("name"), "kingdom": h.get("kingdom"),
                 "reads": h.get("reads_display"), "comment": h.get("comment")}
                for h in top_hits
            ],
        }

    summary = {
        "sample": args.sample,
        "platform": args.platform,
        "date": date.today().isoformat(),
        "hg19_alignment": host,
        "ichorcna": {
            "panel_of_normals": "hg19 (nanoDx-Parameter, r-ichorcna 0.5.1)",
            **{k: v for k, v in cnv.items() if k != "plot_path"},
        },
        "krakenuniq": {
            "enabled": krakenuniq_ran,
            "input": "all reads",
            **(krakenuniq_json or {}),
        },
    }
    with open(args.output_json, "w") as fh:
        json.dump(summary, fh, indent=2)
    print(f"Summary JSON -> {args.output_json}")


if __name__ == "__main__":
    main()
